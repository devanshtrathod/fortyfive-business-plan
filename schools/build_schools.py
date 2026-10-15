"""Build the FortyFive school target list from the scraped IB directory pages.

Inputs : tools/cache/list_*.html and tools/cache/s_*.html (from tools/scrape_ib.py + tools/scrape_pw.py)
Outputs: FortyFive School Targets.xlsx  (the working tracker: edit Status / Next step / Notes in place)
         schools.json                    (normalised rows, embedded in the dashboard)
         summary.json                    (counts used by the deck and the dashboard)
Run: python build_schools.py [path-to-cache]
Re-running rebuilds the xlsx; copy your Status/Notes columns across first if you have edited it (the script
keeps them automatically when the file already exists, matched on IB school code).
"""
import glob
import json
import os
import re
import sys
from collections import Counter, defaultdict

HERE = os.path.dirname(os.path.abspath(__file__))
TOOLS = os.path.join(HERE, 'tools')
CACHE = sys.argv[1] if len(sys.argv) > 1 else os.path.join(TOOLS, 'cache')
sys.path.insert(0, TOOLS)
import scrape_ib as S  # noqa: E402

COUNTRIES = ['IN', 'AE', 'QA', 'SA', 'OM', 'KW', 'BH', 'NP', 'LK', 'BD', 'SG', 'MY', 'ID', 'TH', 'HK', 'KE', 'TZ', 'UG', 'ZA', 'AU', 'NZ']
CNAME = {'IN': 'India', 'AE': 'UAE', 'QA': 'Qatar', 'SA': 'Saudi Arabia', 'OM': 'Oman', 'KW': 'Kuwait', 'BH': 'Bahrain',
         'NP': 'Nepal', 'LK': 'Sri Lanka', 'BD': 'Bangladesh', 'SG': 'Singapore', 'MY': 'Malaysia', 'ID': 'Indonesia',
         'TH': 'Thailand', 'HK': 'Hong Kong', 'KE': 'Kenya', 'TZ': 'Tanzania', 'UG': 'Uganda', 'ZA': 'South Africa',
         'AU': 'Australia', 'NZ': 'New Zealand'}
GROUP = {'IN': 'India', 'AE': 'Gulf', 'QA': 'Gulf', 'SA': 'Gulf', 'OM': 'Gulf', 'KW': 'Gulf', 'BH': 'Gulf',
         'NP': 'South Asia', 'LK': 'South Asia', 'BD': 'South Asia', 'SG': 'South-east and east Asia', 'MY': 'South-east and east Asia',
         'ID': 'South-east and east Asia', 'TH': 'South-east and east Asia', 'HK': 'South-east and east Asia',
         'KE': 'Africa', 'TZ': 'Africa', 'UG': 'Africa', 'ZA': 'Africa', 'AU': 'Oceania', 'NZ': 'Oceania'}
GROUP_SCORE = {'Gulf': 40, 'South Asia': 30, 'South-east and east Asia': 25, 'Africa': 15, 'Oceania': 15}

CITY_ALIASES = {
    'Navi Mumbai': 'Navi Mumbai', 'Mumbai': 'Mumbai', 'Bombay': 'Mumbai', 'Thane': 'Thane', 'Pune': 'Pune', 'Bengaluru': 'Bengaluru',
    'Bangalore': 'Bengaluru', 'Chennai': 'Chennai', 'Hyderabad': 'Hyderabad', 'Secunderabad': 'Hyderabad', 'Coimbatore': 'Coimbatore',
    'Kochi': 'Kochi', 'Cochin': 'Kochi', 'Ernakulam': 'Kochi', 'Thiruvananthapuram': 'Thiruvananthapuram', 'Trivandrum': 'Thiruvananthapuram',
    'Kozhikode': 'Kozhikode', 'Calicut': 'Kozhikode', 'Thrissur': 'Thrissur', 'Mangaluru': 'Mangaluru', 'Mangalore': 'Mangaluru',
    'Mysuru': 'Mysuru', 'Mysore': 'Mysuru', 'Udaipur': 'Udaipur', 'Jaipur': 'Jaipur', 'Jodhpur': 'Jodhpur', 'New Delhi': 'Delhi',
    'Delhi': 'Delhi', 'Gurugram': 'Gurugram', 'Gurgaon': 'Gurugram', 'Greater Noida': 'Greater Noida', 'Noida': 'Noida',
    'Faridabad': 'Faridabad', 'Ghaziabad': 'Ghaziabad', 'Chandigarh': 'Chandigarh', 'Mohali': 'Mohali', 'Panchkula': 'Panchkula',
    'Ludhiana': 'Ludhiana', 'Amritsar': 'Amritsar', 'Jalandhar': 'Jalandhar', 'Dehradun': 'Dehradun', 'Mussoorie': 'Mussoorie',
    'Lucknow': 'Lucknow', 'Kanpur': 'Kanpur', 'Varanasi': 'Varanasi', 'Ahmedabad': 'Ahmedabad', 'Gandhinagar': 'Gandhinagar',
    'Surat': 'Surat', 'Vadodara': 'Vadodara', 'Baroda': 'Vadodara', 'Rajkot': 'Rajkot', 'Indore': 'Indore', 'Bhopal': 'Bhopal',
    'Nagpur': 'Nagpur', 'Nashik': 'Nashik', 'Aurangabad': 'Aurangabad', 'Kolhapur': 'Kolhapur', 'Panaji': 'Goa', 'Goa': 'Goa',
    'Kolkata': 'Kolkata', 'Calcutta': 'Kolkata', 'Bhubaneswar': 'Bhubaneswar', 'Patna': 'Patna', 'Ranchi': 'Ranchi', 'Guwahati': 'Guwahati',
    'Visakhapatnam': 'Visakhapatnam', 'Vijayawada': 'Vijayawada', 'Kodaikanal': 'Kodaikanal', 'Ooty': 'Ooty', 'Udhagamandalam': 'Ooty',
    'Madurai': 'Madurai', 'Tiruchirappalli': 'Tiruchirappalli', 'Trichy': 'Tiruchirappalli', 'Salem': 'Salem', 'Vellore': 'Vellore',
    'Puducherry': 'Puducherry', 'Pondicherry': 'Puducherry', 'Sonipat': 'Sonipat', 'Raipur': 'Raipur', 'Jammu': 'Jammu',
    'Srinagar': 'Srinagar', 'Shimla': 'Shimla', 'Siliguri': 'Siliguri', 'Darjeeling': 'Darjeeling', 'Lonavala': 'Lonavala',
    'Panchgani': 'Panchgani', 'Dharamshala': 'Dharamshala', 'Neemrana': 'Neemrana', 'Ajmer': 'Ajmer', 'Bhiwadi': 'Bhiwadi',
    'Hosur': 'Hosur', 'Tiruppur': 'Tiruppur', 'Erode': 'Erode', 'Palakkad': 'Palakkad', 'Kottayam': 'Kottayam', 'Kollam': 'Kollam',
    'Belagavi': 'Belagavi', 'Belgaum': 'Belagavi', 'Hubli': 'Hubballi', 'Hubballi': 'Hubballi', 'Gwalior': 'Gwalior', 'Agra': 'Agra',
    'Meerut': 'Meerut', 'Haridwar': 'Haridwar', 'Nainital': 'Nainital', 'Bhuj': 'Bhuj', 'Anand': 'Anand', 'Vapi': 'Vapi',
    'Tirupati': 'Tirupati', 'Guntur': 'Guntur', 'Nellore': 'Nellore', 'Warangal': 'Warangal', 'Durgapur': 'Durgapur',
    'Jamshedpur': 'Jamshedpur', 'Cuttack': 'Cuttack', 'Shillong': 'Shillong', 'Gangtok': 'Gangtok', 'Alibaug': 'Alibaug',
    'Karjat': 'Karjat', 'Mahabaleshwar': 'Mahabaleshwar', 'Jalgaon': 'Jalgaon', 'Solapur': 'Solapur', 'Sangli': 'Sangli',
    'Satara': 'Satara', 'Bathinda': 'Bathinda', 'Patiala': 'Patiala', 'Ambala': 'Ambala', 'Karnal': 'Karnal', 'Rohtak': 'Rohtak',
    'Hisar': 'Hisar', 'Rewari': 'Rewari', 'Jabalpur': 'Jabalpur', 'Ujjain': 'Ujjain', 'Kota': 'Kota', 'Bikaner': 'Bikaner',
    'Allahabad': 'Prayagraj', 'Prayagraj': 'Prayagraj', 'Bareilly': 'Bareilly', 'Gorakhpur': 'Gorakhpur', 'Aligarh': 'Aligarh',
}
INTL_CITIES = {c: c for c in (
    'Doha', 'Lusail', 'Al Khor', 'Al Wakrah', 'Dubai', 'Abu Dhabi', 'Sharjah', 'Ajman', 'Al Ain', 'Ras Al Khaimah', 'Fujairah',
    'Riyadh', 'Jeddah', 'Dammam', 'Al Khobar', 'Dhahran', 'Jubail', 'Mecca', 'Medina', 'Thuwal', 'Muscat', 'Sohar', 'Salalah',
    'Kuwait City', 'Salmiya', 'Manama', 'Isa Town', 'Saar', 'Riffa', 'Kathmandu', 'Lalitpur', 'Colombo', 'Dhaka', 'Chittagong',
    'Singapore', 'Kuala Lumpur', 'Petaling Jaya', 'Penang', 'Johor Bahru', 'Iskandar Puteri', 'Kota Kinabalu', 'Jakarta', 'Tangerang',
    'Bekasi', 'Bogor', 'Bandung', 'Surabaya', 'Semarang', 'Medan', 'Bali', 'Denpasar', 'Badung', 'Yogyakarta', 'Malang', 'Batam',
    'Balikpapan', 'Makassar', 'Bangkok', 'Nonthaburi', 'Pathum Thani', 'Chiang Mai', 'Phuket', 'Pattaya', 'Chon Buri', 'Hua Hin',
    'Hong Kong', 'Kowloon', 'Nairobi', 'Mombasa', 'Kisumu', 'Dar es Salaam', 'Arusha', 'Moshi', 'Kampala', 'Entebbe', 'Johannesburg',
    'Cape Town', 'Pretoria', 'Durban', 'Sydney', 'Melbourne', 'Brisbane', 'Perth', 'Adelaide', 'Canberra', 'Hobart', 'Darwin',
    'Gold Coast', 'Geelong', 'Auckland', 'Wellington', 'Christchurch', 'Hamilton', 'Dunedin', 'Tauranga', 'Nelson')}
INTL_CITIES.update({'Jakarta Selatan': 'Jakarta', 'South Jakarta': 'Jakarta', 'Tangerang Selatan': 'Tangerang', 'Ar Rayyan': 'Doha',
                    'Al Rayyan': 'Doha', 'Al-Wukair': 'Doha', 'Kowloon Tong': 'Hong Kong', 'Nonthaburi Province': 'Nonthaburi'})
PIN2STATE = [((110, 110), 'Delhi'), ((120, 136), 'Haryana'), ((140, 159), 'Punjab'), ((160, 160), 'Chandigarh'),
             ((161, 169), 'Punjab'), ((170, 177), 'Himachal Pradesh'), ((180, 194), 'Jammu and Kashmir'), ((246, 249), 'Uttarakhand'),
             ((262, 263), 'Uttarakhand'), ((200, 285), 'Uttar Pradesh'), ((300, 345), 'Rajasthan'), ((360, 396), 'Gujarat'),
             ((403, 403), 'Goa'), ((400, 445), 'Maharashtra'), ((450, 488), 'Madhya Pradesh'), ((490, 497), 'Chhattisgarh'),
             ((500, 509), 'Telangana'), ((510, 535), 'Andhra Pradesh'), ((560, 591), 'Karnataka'), ((605, 605), 'Puducherry'),
             ((600, 643), 'Tamil Nadu'), ((670, 695), 'Kerala'), ((737, 737), 'Sikkim'), ((700, 743), 'West Bengal'),
             ((744, 744), 'Andaman and Nicobar'), ((750, 770), 'Odisha'), ((780, 788), 'Assam'), ((790, 792), 'Arunachal Pradesh'),
             ((793, 794), 'Meghalaya'), ((795, 795), 'Manipur'), ((796, 796), 'Mizoram'), ((797, 798), 'Nagaland'), ((799, 799), 'Tripura'),
             ((814, 835), 'Jharkhand'), ((800, 855), 'Bihar')]
INDIAN_LINK = re.compile(r'\b(indian|india|giis|global indian|delhi|bhavan|modern academy|dps)\b', re.I)
SUBJ = {'English A L&L': r'ENGLISH A LAL', 'French ab initio': r'FRENCH AB', 'Maths AI': r'MATH APP', 'Maths AA': r'MATH ANAL',
        'Business': r'BUS MAN', 'Economics': r'ECONOMICS', 'Computer science': r'COMPUTER SC'}


def pin_state(pin):
    p3 = int(pin[:3])
    for (a, b), st in PIN2STATE:
        if a <= p3 <= b:
            return st
    return ''


def city_from(addr, cc):
    lines = [x for x in addr if x and x.upper() not in ('INDIA', CNAME.get(cc, '').upper())]
    pin_i = next((i for i, x in enumerate(lines) if re.fullmatch(r'\d{6}', x.replace(' ', ''))), None)
    names = sorted(CITY_ALIASES, key=len, reverse=True)
    order = list(range(pin_i - 1, -1, -1)) if pin_i is not None else list(range(len(lines) - 1, -1, -1))
    if cc == 'IN':
        for i in order + list(range(len(lines))):
            for n in names:
                if re.search(r'(?i)(?<![a-z])' + re.escape(n), lines[i]) or n.lower() in lines[i].lower():
                    return CITY_ALIASES[n]
    elif cc in ('SG', 'HK'):
        return CNAME[cc]
    else:
        joined = ' | '.join(lines)
        for n in sorted(INTL_CITIES, key=len, reverse=True):
            if re.search(r'(?i)\b' + re.escape(n) + r'\b', joined):
                return INTL_CITIES[n]
        # otherwise the last line that looks like a place: no digits, not a PO box, not the country
        for x in reversed(lines):
            if not re.search(r'\d|p\.?\s?o\.?\s?box|street|road|building|zone|c/o|\(', x, flags=re.I) and len(x) > 2 and x.upper() != CNAME.get(cc, '').upper():
                return x.title() if (x.isupper() and len(x) > 3) else x
    # no match: the line before the postcode / the last line, tidied
    cand = lines[pin_i - 1] if pin_i else (lines[-1] if lines else '')
    cand = re.sub(r'\b\d{3,}\b', '', cand).strip(' ,-')
    return cand.title() if cand.isupper() else cand


def load():
    rows, seen = [], set()
    for cc in COUNTRIES:
        for f in sorted(glob.glob(os.path.join(CACHE, 'list_%s_*.html' % cc))):
            _, rs = S.parse_list(open(f, encoding='utf-8', errors='ignore').read())
            for r in rs:
                if r['code'] in seen:
                    continue
                seen.add(r['code'])
                p = os.path.join(CACHE, 's_%s.html' % r['code'])
                d = {}
                if os.path.exists(p):
                    h = open(p, encoding='utf-8', errors='ignore').read()
                    if 'PropertyList-key' in h:
                        d = S.parse_school(h, r['code'])
                d['code'] = r['code']
                d['cc'] = cc
                d['name'] = d.get('name') or r['name']
                d['list_languages'] = r['languages']
                d['fetched'] = bool(d.get('dp') is not None and d.get('type') is not None)
                rows.append(d)
    return rows


def normalise(d):
    cc = d['cc']
    dp = d.get('dp') or {}
    subs = dp.get('subjects', [])
    addr = d.get('address', [])
    pin = next((x.replace(' ', '') for x in addr if re.fullmatch(r'\d{6}', x.replace(' ', ''))), '')
    exams = (dp.get('exams') or '').upper()
    sess = 'Nov' if 'NOV' in exams else ('May' if 'MAY' in exams else ('May and Nov' if 'MAY' in exams and 'NOV' in exams else ''))
    if 'MAY' in exams and 'NOV' in exams:
        sess = 'May and Nov'
    since = re.search(r'(\d{4})', dp.get('authorised', '') or '')
    o = {
        'code': d['code'], 'name': d['name'], 'cc': cc, 'country': CNAME.get(cc, cc), 'group': GROUP.get(cc, ''),
        'city': city_from(addr, cc) if addr else '', 'state': pin_state(pin) if (cc == 'IN' and pin) else '',
        'session': sess, 'dp_since': int(since.group(1)) if since else None, 'type': (d.get('type') or '').title(),
        'boarding': (dp.get('boarding') or '').title(), 'language': (dp.get('language') or ', '.join(d.get('list_languages', []))).title(),
        'n_subjects': len(subs), 'subjects': ', '.join(s.title() for s in subs), 'website': d.get('website', ''),
        'coordinator': d.get('coordinator', ''), 'ib_page': 'https://www.ibo.org/school/%s/' % d['code'],
        'km': d.get('km_from_coimbatore'), 'fetched': d['fetched'],
    }
    for k, pat in SUBJ.items():
        o[k] = any(re.search(pat, s) for s in subs)
    o['indian_link'] = bool(cc != 'IN' and INDIAN_LINK.search(d['name']))
    # tiers and scores
    if cc == 'IN':
        km = o['km'] if o['km'] is not None else 9999
        o['tier'] = 'Tier 1: within 450 km' if km <= 450 else ('Tier 2: 450 to 1,100 km' if km <= 1100 else 'Tier 3: further')
        base = 50 if km <= 450 else (25 if km <= 1100 else 10)
    else:
        o['tier'] = o['group']
        base = GROUP_SCORE.get(o['group'], 10)
        if 'English' not in o['language'] and o['language']:
            base -= 20
        if o['indian_link']:
            base += 10
    score = base + (20 if 'Nov' in o['session'] else 0) + (10 if o['English A L&L'] else 0) + (8 if o['French ab initio'] else 0)
    score += min(o['n_subjects'], 30) / 3
    o['score'] = round(score, 1)
    return o


def main():
    raw = load()
    rows = [normalise(d) for d in raw]
    india = sorted([r for r in rows if r['cc'] == 'IN'], key=lambda r: -r['score'])
    intl = sorted([r for r in rows if r['cc'] != 'IN'], key=lambda r: -r['score'])
    for lst in (india, intl):
        for i, r in enumerate(lst, 1):
            r['rank'] = i
            r['priority'] = 'A' if i <= (30 if lst is india else 40) else ('B' if i <= (90 if lst is india else 140) else 'C')
    # chains: brand words shared by 2+ schools
    brand = defaultdict(list)
    stop = set('the international school schools academy world of and college high public campus global american australian british '
               'canadian singapore united sekolah emirates indian india english french german japanese korean swiss european asian pacific '
               'new st saint modern dubai doha qatar jakarta bangkok hong kong nairobi kuala lumpur riyadh jeddah muscat kuwait bahrain '
               'abu dhabi sharjah grammar primary secondary centre center institute de la el al junior senior boys girls private community '
               'christian islamic catholic day boarding education group foundation trust city park road branch campus north south east west '
               'national king kings london deutsche german'.split())
    for r in rows:
        low = r['name'].lower()
        if 'global indian' in low or re.search(r'\bgiis\b', low):
            brand['giis'].append(r)
            continue
        words = [w for w in re.findall(r'[a-z]+', low) if w not in stop and len(w) > 2]
        if words:
            brand[' '.join(words[:2]) if len(words) > 1 and len(words[0]) < 4 else words[0]].append(r)
    chains = sorted([(k, v) for k, v in brand.items() if len(v) >= 2 and len(k) > 3], key=lambda kv: -len(kv[1]))
    by_country = Counter(r['cc'] for r in rows)
    nov = Counter(r['cc'] for r in rows if 'Nov' in r['session'])
    summary = {
        'in_dp': by_country['IN'], 'in_nov': nov['IN'], 'intl_dp': sum(v for k, v in by_country.items() if k != 'IN'),
        'intl_nov': sum(v for k, v in nov.items() if k != 'IN'),
        'fetched': sum(1 for r in rows if r['fetched']), 'total': len(rows),
        'in_tier1': sum(1 for r in india if r['tier'].startswith('Tier 1')), 'in_tier2': sum(1 for r in india if r['tier'].startswith('Tier 2')),
        'in_tier3': sum(1 for r in india if r['tier'].startswith('Tier 3')),
        'in_lal': sum(1 for r in india if r['English A L&L']), 'in_french': sum(1 for r in india if r['French ab initio']),
        'by_country': [{'cc': k, 'name': CNAME[k], 'dp': by_country[k], 'nov': nov[k]} for k in sorted(by_country, key=lambda k: -by_country[k])],
        'by_group': [{'group': g, 'dp': sum(1 for r in rows if r['group'] == g), 'nov': sum(1 for r in rows if r['group'] == g and 'Nov' in r['session'])}
                     for g in ['India', 'Gulf', 'South Asia', 'South-east and east Asia', 'Africa', 'Oceania']],
        'in_cities': Counter(r['city'] or 'Unknown' for r in india).most_common(20),
        'chains': [{'brand': k.title(), 'n': len(v), 'countries': sorted({x['country'] for x in v}), 'schools': [x['name'] for x in v]} for k, v in chains[:25]],
    }
    json.dump(rows, open(os.path.join(HERE, 'schools.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=0)
    json.dump(summary, open(os.path.join(HERE, 'summary.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    # the dashboard loads this as a plain script (works from file://); no coordinator names in it
    keys = ('code', 'name', 'city', 'state', 'country', 'cc', 'group', 'session', 'tier', 'priority', 'rank', 'score', 'km',
            'website', 'ib_page', 'n_subjects', 'dp_since', 'indian_link', 'fetched')
    mini = [dict({k: r[k] for k in keys}, lal=r['English A L&L'], fr=r['French ab initio']) for r in india + intl]
    with open(os.path.join(HERE, 'schools_data.js'), 'w', encoding='utf-8') as f:
        f.write('/* generated by build_schools.py from the IB school directory; do not edit by hand */\n')
        f.write('window.FF_SCHOOLS = ' + json.dumps(mini, ensure_ascii=False) + ';\n')
        f.write('window.FF_SUMMARY = ' + json.dumps(summary, ensure_ascii=False) + ';\n')
    write_xlsx(india, intl, summary)
    print('schools', len(rows), 'fetched', summary['fetched'], '| India', summary['in_dp'], 'Nov', summary['in_nov'],
          '| intl', summary['intl_dp'], 'Nov', summary['intl_nov'])


def write_xlsx(india, intl, summary):
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    path = os.path.join(HERE, 'FortyFive School Targets.xlsx')
    keep = {}
    if os.path.exists(path):  # carry over anything typed into the tracking columns
        old = load_workbook(path)
        for ws in old.worksheets:
            hdr = [c.value for c in ws[1]]
            if 'IB code' in hdr and 'Status' in hdr:
                ci, cols = hdr.index('IB code'), [h for h in ('Status', 'Next step', 'Last contact', 'Notes') if h in hdr]
                for row in ws.iter_rows(min_row=2, values_only=True):
                    if row[ci]:
                        keep[str(row[ci])] = {h: row[hdr.index(h)] for h in cols}
    wb = Workbook()
    head_fill = PatternFill('solid', fgColor='0B0B0B')
    head_font = Font(bold=True, color='FFFFFF')
    nov_fill = PatternFill('solid', fgColor='E9F1FB')
    statuses = '"Not contacted,Emailed,Replied,Demo booked,Pilot,Paying,Not now"'

    def sheet(ws, rows, cols):
        ws.append([c[0] for c in cols])
        for r in rows:
            vals = []
            for title, key, _w in cols:
                if key in ('Status', 'Next step', 'Last contact', 'Notes'):
                    vals.append(keep.get(r['code'], {}).get(key) or ('Not contacted' if key == 'Status' else None))
                elif key in SUBJ or key == 'indian_link':
                    vals.append('Yes' if r[key] else '')
                else:
                    vals.append(r.get(key))
            ws.append(vals)
        for i, (title, key, w) in enumerate(cols, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
            c = ws.cell(row=1, column=i)
            c.fill, c.font = head_fill, head_font
            c.alignment = Alignment(vertical='center', wrap_text=True)
        ws.row_dimensions[1].height = 30
        ws.freeze_panes = 'D2'
        ws.auto_filter.ref = 'A1:%s%d' % (get_column_letter(len(cols)), len(rows) + 1)
        s_col = [i for i, c in enumerate(cols, 1) if c[1] == 'session'][0]
        for row in ws.iter_rows(min_row=2):
            if row[s_col - 1].value and 'Nov' in str(row[s_col - 1].value):
                for c in row[:4]:
                    c.fill = nov_fill
            for c in row:
                if c.column_letter and isinstance(c.value, str) and c.value.startswith('http'):
                    c.hyperlink = c.value
                    c.font = Font(color='2A78D6', underline='single')
        st_col = get_column_letter([i for i, c in enumerate(cols, 1) if c[1] == 'Status'][0])
        dv = DataValidation(type='list', formula1=statuses, allow_blank=True)
        ws.add_data_validation(dv)
        dv.add('%s2:%s%d' % (st_col, st_col, len(rows) + 1))

    common_tail = [('Website', 'website', 34), ('DP coordinator (as listed by the IB)', 'coordinator', 26), ('IB page', 'ib_page', 34),
                   ('IB code', 'code', 9), ('Status', 'Status', 15), ('Next step', 'Next step', 26), ('Last contact', 'Last contact', 13),
                   ('Notes', 'Notes', 40), ('Subjects registered', 'subjects', 60)]
    ws = wb.active
    ws.title = 'India'
    sheet(ws, india, [('Priority', 'priority', 8), ('Rank', 'rank', 6), ('Score', 'score', 7), ('School', 'name', 42),
                      ('City', 'city', 16), ('State', 'state', 15), ('Tier', 'tier', 22), ('Km from Coimbatore', 'km', 10),
                      ('Exams', 'session', 9), ('DP since', 'dp_since', 9), ('Type', 'type', 9), ('Boarding', 'boarding', 11),
                      ('Subjects', 'n_subjects', 9), ('English A L&L', 'English A L&L', 9), ('French ab initio', 'French ab initio', 9),
                      ('Maths AI', 'Maths AI', 8), ('Maths AA', 'Maths AA', 8), ('Business', 'Business', 9), ('Economics', 'Economics', 9),
                      ('Computer science', 'Computer science', 9)] + common_tail)
    ws2 = wb.create_sheet('International')
    sheet(ws2, intl, [('Priority', 'priority', 8), ('Rank', 'rank', 6), ('Score', 'score', 7), ('School', 'name', 42),
                      ('City', 'city', 16), ('Country', 'country', 13), ('Market', 'group', 20), ('Exams', 'session', 9),
                      ('DP since', 'dp_since', 9), ('Language', 'language', 14), ('Indian link (check)', 'indian_link', 10),
                      ('Subjects', 'n_subjects', 9), ('English A L&L', 'English A L&L', 9), ('French ab initio', 'French ab initio', 9),
                      ('Maths AI', 'Maths AI', 8), ('Business', 'Business', 9), ('Economics', 'Economics', 9)] + common_tail)
    ws3 = wb.create_sheet('Summary')
    ws3.append(['Country', 'Diploma schools', 'November session'])
    for c in summary['by_country']:
        ws3.append([c['name'], c['dp'], c['nov']])
    ws3.append([])
    ws3.append(['Market', 'Diploma schools', 'November session'])
    for g in summary['by_group']:
        ws3.append([g['group'], g['dp'], g['nov']])
    ws3.append([])
    ws3.append(['School groups (same brand, 2+ schools)', 'Schools', 'Countries'])
    for ch in summary['chains']:
        ws3.append([ch['brand'], ch['n'], ', '.join(ch['countries'])])
    for col, w in (('A', 40), ('B', 16), ('C', 40)):
        ws3.column_dimensions[col].width = w
    for r in (1,):
        for c in ws3[r]:
            c.font = Font(bold=True)
    ws4 = wb.create_sheet('How to use')
    notes = [
        'FortyFive school targets, built from the IB "Find an IB World School" directory (October 2026).',
        '',
        'Priority A/B/C comes from Score. India: distance from Coimbatore (Tier 1 within 450 km), November session (+20),',
        'English A Language and Literature (+10), French ab initio (+8), and number of DP subjects as a size signal.',
        'International: market (Gulf first), November session, English-medium, likely Indian community school, subjects.',
        'Blue rows sit the November session, like you: your predicted November mocks fit them first.',
        '',
        'How to work it:',
        '1. Nothing before 13 November 2026.',
        '2. Start with Priority A in Tier 1. Contact the DP coordinator through the school\'s official email or the',
        '   "Contact coordinator" form on the IB page. Never contact students directly.',
        '3. Offer the free pilot (February to March 2027) and a 15-minute demo. Use the email in the dashboard\'s Scripts tab.',
        '4. Move Status along: Not contacted, Emailed, Replied, Demo booked, Pilot, Paying, Not now. Log the date.',
        '5. Re-run build_schools.py after a fresh scrape; your Status, Next step, Last contact and Notes are kept.',
        '',
        'Check before relying on a row: cities are read from the IB address and can be wrong; "Indian link" is a guess from the name.',
    ]
    for n in notes:
        ws4.append([n])
    ws4.column_dimensions['A'].width = 120
    wb.move_sheet('How to use', offset=-3)
    wb.save(path)


if __name__ == '__main__':
    main()
